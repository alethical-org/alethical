#!/usr/bin/env python3
"""Build the manually refreshed state reference from HUD's ZIP-to-county file.

Download the full national ZIP-to-county workbook from HUD User after sign-in.
Then run this command in an isolated worktree and review its diff through a PR:

    uv run python scripts/build_zip_state_reference.py \
        --input /path/to/ZIP_COUNTY.xlsx --as-of 2026-06-30 \
        --copied-at 2026-09-13T00:00:00+00:00 \
        --output alethical/api/data/zip_states.json

The dates above are examples, not defaults. Preserve the actual quarter-end and
copy time. No network request, database write, donor lookup or scheduled refresh
runs here. A ZIP crossing state boundaries remains unknown. The mailing state's
preferred name and the address ratios do not assign an individual donor a state.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
import csv
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from alethical.api.services.zip_state_reference import SUPPORTED_STATES  # noqa: E402

SOURCE_URL = "https://www.huduser.gov/portal/datasets/usps_crosswalk.html"
# Census's authoritative state portion of the 5-digit county GEOID:
# https://www.census.gov/library/reference/code-lists/ansi/ansi-codes-for-states.html
STATE_CODES = dict(
    pair.split(":")
    for pair in (
        "01:AL 02:AK 04:AZ 05:AR 06:CA 08:CO 09:CT 10:DE 11:DC 12:FL "
        "13:GA 15:HI 16:ID 17:IL 18:IN 19:IA 20:KS 21:KY 22:LA 23:ME "
        "24:MD 25:MA 26:MI 27:MN 28:MS 29:MO 30:MT 31:NE 32:NV 33:NH "
        "34:NJ 35:NM 36:NY 37:NC 38:ND 39:OH 40:OK 41:OR 42:PA 44:RI "
        "45:SC 46:SD 47:TN 48:TX 49:UT 50:VT 51:VA 53:WA 54:WV 55:WI "
        "56:WY 60:AS 64:FM 66:GU 68:MH 69:MP 70:PW 72:PR 74:UM 78:VI"
    ).split()
)


def source_code(value: object) -> str:
    """HUD defines both columns as full 5-digit codes, including Excel numerics.

    This conversion is for the reference workbook only. Donor ZIPs are text as
    filed, and the public service never fills in their missing leading zeros.
    """
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and 0 <= value < 100000
        and value == int(value)
    ):
        return f"{int(value):05d}"
    if isinstance(value, str) and re.fullmatch(r"[0-9]{5}", value.strip()):
        return value.strip()
    raise ValueError("The HUD reference contains an unreadable 5-digit code")


def parse_rows(rows) -> tuple[dict[str, str | None], int]:
    iterator = iter(rows)
    headers = tuple(str(value).strip().upper() for value in next(iterator, ()))
    if len(set(headers)) != len(headers) or not {"ZIP", "COUNTY"} <= set(headers):
        raise ValueError("Expected HUD ZIP-to-county columns ZIP and COUNTY")
    zip_index, county_index = headers.index("ZIP"), headers.index("COUNTY")
    states: dict[str, set[str]] = defaultdict(set)
    count = 0
    for values in iterator:
        if not any(value not in (None, "") for value in values):
            continue
        if len(values) != len(headers):
            raise ValueError("A HUD reference row has a different number of columns")
        zipcode, county = (
            source_code(values[zip_index]),
            source_code(values[county_index]),
        )
        state = STATE_CODES.get(county[:2])
        if state is None:
            raise ValueError("The HUD county code has an unrecognized state")
        states[zipcode].add(state)
        count += 1
    if not count:
        raise ValueError("The HUD reference contains no ZIP-to-county rows")
    return {
        zipcode: next(iter(found))
        if len(found) == 1 and next(iter(found)) in SUPPORTED_STATES
        else None
        for zipcode, found in sorted(states.items())
    }, count


def read_reference(path: Path) -> tuple[dict[str, str | None], int]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return parse_rows(csv.reader(handle))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Use the HUD .xlsx workbook or a text-preserving .csv export")
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected 1 sheet in the HUD ZIP-to-county workbook")
        return parse_rows(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--as-of", type=date.fromisoformat, required=True)
    parser.add_argument("--copied-at", type=datetime.fromisoformat, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.copied_at.tzinfo is None:
        parser.error("--copied-at needs its UTC offset")
    if (args.as_of.month, args.as_of.day) not in ((3, 31), (6, 30), (9, 30), (12, 31)):
        parser.error("--as-of must be the source quarter's last day")
    if args.as_of > args.copied_at.date():
        parser.error("The source quarter cannot end after the file was copied")
    states, rows = read_reference(args.input)
    # Reject a state-only extract. Seeing every state plus DC cannot prove every
    # county row is present: the complete authenticated source file, its hash,
    # and the real allocation check must still be reviewed before release.
    if not SUPPORTED_STATES <= set(states.values()):
        parser.error("The national reference must include all 50 states and DC")
    value = {
        "source_url": SOURCE_URL,
        "as_of": args.as_of.isoformat(),
        "copied_at": args.copied_at.isoformat(),
        "content_hash": hashlib.sha256(args.input.read_bytes()).hexdigest(),
        "states": states,
    }
    args.output.write_text(json.dumps(value, indent=2) + "\n")
    print(
        json.dumps(
            {
                "source_rows": rows,
                "distinct_zips": len(states),
                "unassigned_zips": sum(state is None for state in states.values()),
                "reference_date": value["as_of"],
                "source_sha256": value["content_hash"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
